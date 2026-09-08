import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # Set slide dimensions to widescreen 16:9 (13.33 x 7.5 inches)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Colors
    CRIMSON = RGBColor(220, 20, 60)       # #DC143C
    ROSE_BG = RGBColor(255, 241, 242)     # #FFF1F2
    DARK_SLATE = RGBColor(15, 23, 42)     # #0F172A
    MUTED_TEXT = RGBColor(100, 116, 139)  # #64748B
    WHITE = RGBColor(255, 255, 255)
    LIGHT_GRAY = RGBColor(241, 245, 249)
    CARD_BORDER = RGBColor(254, 205, 211)

    base_dir = r"C:\Users\rajas\.gemini\antigravity\brain\1a3d3ab8-b9e4-4d1c-a382-11dc322a4a0d\.user_uploaded"
    img_landing = os.path.join(base_dir, "media__1788264042297.png")
    img_caregiver = os.path.join(base_dir, "media__1788264042302.png")
    img_patient = os.path.join(base_dir, "media__1788264042305.png")
    img_admin = os.path.join(base_dir, "media__1788264042382.png")

    blank_slide_layout = prs.slide_layouts[6]

    def add_header(slide, title_text, category_text="PillSync Milestone 1"):
        # Header background banner
        header_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.1))
        header_shape.fill.solid()
        header_shape.fill.fore_color.rgb = ROSE_BG
        header_shape.line.color.rgb = CARD_BORDER

        # Category pill
        txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.12), Inches(5), Inches(0.3))
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        p.text = category_text.upper()
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = CRIMSON

        # Main Title
        txBox2 = slide.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(10), Inches(0.6))
        tf2 = txBox2.text_frame
        p2 = tf2.paragraphs[0]
        p2.text = title_text
        p2.font.size = Pt(22)
        p2.font.bold = True
        p2.font.color.rgb = DARK_SLATE

    # ==========================================
    # SLIDE 1: Title Slide
    # ==========================================
    slide1 = prs.slides.add_slide(blank_slide_layout)
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = ROSE_BG
    bg1.line.fill.background()

    card1 = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.2), Inches(10.333), Inches(5.1))
    card1.fill.solid()
    card1.fill.fore_color.rgb = WHITE
    card1.line.color.rgb = CARD_BORDER

    tb1 = slide1.shapes.add_textbox(Inches(2.0), Inches(1.8), Inches(9.333), Inches(1.0))
    p = tb1.text_frame.paragraphs[0]
    p.text = "PillSync"
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = CRIMSON
    p.alignment = PP_ALIGN.CENTER

    tb2 = slide1.shapes.add_textbox(Inches(2.0), Inches(2.7), Inches(9.333), Inches(0.8))
    p = tb2.text_frame.paragraphs[0]
    p.text = "Intelligent Medicine Reminder & Tracking Platform"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = DARK_SLATE
    p.alignment = PP_ALIGN.CENTER

    tb3 = slide1.shapes.add_textbox(Inches(2.0), Inches(3.6), Inches(9.333), Inches(0.6))
    p = tb3.text_frame.paragraphs[0]
    p.text = "Milestone 1: Frontend Setup, UI Wireframes & Workflow Planning"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = CRIMSON
    p.alignment = PP_ALIGN.CENTER

    tb4 = slide1.shapes.add_textbox(Inches(2.0), Inches(4.7), Inches(9.333), Inches(1.0))
    tf4 = tb4.text_frame
    p1 = tf4.paragraphs[0]
    p1.text = "Presenter: Rajasri Nallamilli (Intern ID #03)"
    p1.font.size = Pt(15)
    p1.font.bold = True
    p1.font.color.rgb = DARK_SLATE
    p1.alignment = PP_ALIGN.CENTER

    p2 = tf4.add_paragraph()
    p2.text = "Infosys Springboard AI Internship Program • September 2026"
    p2.font.size = Pt(13)
    p2.font.color.rgb = MUTED_TEXT
    p2.alignment = PP_ALIGN.CENTER

    # ==========================================
    # SLIDE 2: Project Overview & Milestone 1 Scope
    # ==========================================
    slide2 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide2, "Project Overview & Milestone 1 Scope", "01. Introduction")

    # Left Box: Problem & Solution
    left_box = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), Inches(5.6), Inches(5.5))
    left_box.fill.solid()
    left_box.fill.fore_color.rgb = WHITE
    left_box.line.color.rgb = CARD_BORDER

    tb_left = slide2.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(5.2), Inches(5.0))
    tf = tb_left.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "🎯 Problem Statement & Goal"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = CRIMSON

    points_left = [
        "• Challenge: Medication non-adherence leads to severe health complications and preventable hospitalizations.",
        "• Solution: PillSync is an intelligent multi-role platform that automates daily dose reminders, stock refill predictions, prescription OCR scanning, and caregiver alerts.",
        "• Internship Scope: Team of 4 members collaborating on full-stack deployment."
    ]
    for pt in points_left:
        p = tf.add_paragraph()
        p.text = pt
        p.font.size = Pt(13)
        p.font.color.rgb = DARK_SLATE

    # Right Box: Assigned Work Boundaries
    right_box = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.4), Inches(5.7), Inches(5.5))
    right_box.fill.solid()
    right_box.fill.fore_color.rgb = ROSE_BG
    right_box.line.color.rgb = CARD_BORDER

    tb_right = slide2.shapes.add_textbox(Inches(7.0), Inches(1.6), Inches(5.3), Inches(5.0))
    tf = tb_right.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "👤 Rajasri's Milestone 1 Deliverables"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = CRIMSON

    points_right = [
        "☑ Frontend Application Setup: Vite, React 18, Custom CSS Tokens (#ED4264, #DC143C).",
        "☑ Complete UI Wireframes Blueprint: 25 screens designed across Patient, Caregiver, and Admin portals.",
        "☑ Workflow Architecture Planning: 4 core sequence & flowchart diagrams.",
        "☑ Role-Based Direct URL Routing: Clean navigation (/patient, /caregiver, /admin).",
        "☐ Handled by Teammates: Backend Init, Auth API (JWT/OAuth), DB Schema, PostgreSQL."
    ]
    for pt in points_right:
        p = tf.add_paragraph()
        p.text = pt
        p.font.size = Pt(13)
        p.font.color.rgb = DARK_SLATE

    # ==========================================
    # SLIDE 3: Design System & Routing Architecture
    # ==========================================
    slide3 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide3, "Design Aesthetics & URL Routing Architecture", "02. Architecture")

    # 3 Cards Layout
    cards_data = [
        ("🎨 Color Palette & Aesthetics", "• Primary Crimson: #ED4264 & #DC143C\n• Soft Rose Background: #FFF1F2\n• Glassmorphic cards with subtle borders.\n• Curved category headers & high contrast text."),
        ("🌐 Direct URL Routing", "• Configured react-router-dom BrowserRouter.\n• Direct URL paths (/patient, /caregiver, /admin).\n• Bypasses OAuth blockers for instant UI testing.\n• Fully responsive across mobile & desktop."),
        ("⚙️ Clean Empty State Design", "• 100% Manual Data Entry controls.\n• Zero hardcoded default dummy data.\n• Interactive Add/Edit forms for Patients, Medicines, Caregivers & Admins.")
    ]

    for i, (ctitle, cdesc) in enumerate(cards_data):
        x = Inches(0.8 + i * 4.0)
        cbox = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.6), Inches(3.7), Inches(5.2))
        cbox.fill.solid()
        cbox.fill.fore_color.rgb = WHITE
        cbox.line.color.rgb = CARD_BORDER

        tb = slide3.shapes.add_textbox(x + Inches(0.15), Inches(1.8), Inches(3.4), Inches(4.8))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = ctitle
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = CRIMSON

        p2 = tf.add_paragraph()
        p2.text = cdesc
        p2.font.size = Pt(13)
        p2.font.color.rgb = DARK_SLATE

    # ==========================================
    # SLIDE 4: Feature 1 - Role Category Selection Hub (Landing Page)
    # ==========================================
    slide4 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide4, "Landing Page — Role Category Selection Hub", "03. UI Wireframes")

    # Left Column: Description
    tb_desc = slide4.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(4.5), Inches(5.5))
    tf = tb_desc.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "🔑 Role Category Selection Hub"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = CRIMSON

    bullets = [
        "• Single Landing Page serving as the central portal entry point.",
        "• Features 3 role category cards with curved top headers:",
        "  1. Patient Portal (Login & /patient direct route)",
        "  2. Caregiver Portal (Login & /caregiver direct route)",
        "  3. Admin Console (Login & /admin direct route)",
        "• High visual impact matching intern reference design requirements.",
        "• Direct access shortcuts in top header bar."
    ]
    for b in bullets:
        p = tf.add_paragraph()
        p.text = b
        p.font.size = Pt(12)
        p.font.color.rgb = DARK_SLATE

    # Right Column: Screenshot
    if os.path.exists(img_landing):
        slide4.shapes.add_picture(img_landing, Inches(5.5), Inches(1.4), width=Inches(7.0))

    # ==========================================
    # SLIDE 5: Feature 2 - Patient Portal Hub
    # ==========================================
    slide5 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide5, "Patient Portal — Dosage Schedule & Health Management", "04. Patient Scope")

    tb_desc = slide5.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(4.5), Inches(5.5))
    tf = tb_desc.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "💊 Patient Dashboard & Profile Hub"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = CRIMSON

    bullets = [
        "• 100% Clean Empty State — No hardcoded dummy medicines.",
        "• Interactive Quick Actions: [ + Add Medicine ] and [ Upload Prescription ].",
        "• Dual Add Method: Manual Entry form & AI/OCR Prescription Image Upload.",
        "• 1-Click Dosage Logger: [ Taken ], [ Missed ], [ Snooze 15m ].",
        "• Consolidated Profile Sub-tabs: Personal Info, Adherence %, Refill Engine, History Logs."
    ]
    for b in bullets:
        p = tf.add_paragraph()
        p.text = b
        p.font.size = Pt(12)
        p.font.color.rgb = DARK_SLATE

    if os.path.exists(img_patient):
        slide5.shapes.add_picture(img_patient, Inches(5.5), Inches(1.4), width=Inches(7.0))

    # ==========================================
    # SLIDE 6: Feature 3 - Caregiver Monitoring Portal
    # ==========================================
    slide6 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide6, "Caregiver Portal — Patient Roster & Alert Hub", "05. Caregiver Scope")

    tb_desc = slide6.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(4.5), Inches(5.5))
    tf = tb_desc.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "👥 Caregiver Dashboard & Alert Feeds"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = CRIMSON

    bullets = [
        "• Monitored Patient Roster: Clean empty state with [ + Assign New Patient ] form.",
        "• Patient Details Page (/caregiver/patients): View compliance, adherence %, and emergency contact info.",
        "• Missed-Dose Alert Stream: Real-time log of missed patient routines.",
        "• Refill Notifications Feed: Stock depletion warnings.",
        "• Caregiver Actions: [ Send Nudge ] & [ Emergency Call Patient ]."
    ]
    for b in bullets:
        p = tf.add_paragraph()
        p.text = b
        p.font.size = Pt(12)
        p.font.color.rgb = DARK_SLATE

    if os.path.exists(img_caregiver):
        slide6.shapes.add_picture(img_caregiver, Inches(5.5), Inches(1.4), width=Inches(7.0))

    # ==========================================
    # SLIDE 7: Feature 4 - Admin Operations Console
    # ==========================================
    slide7 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide7, "Admin Console — User Management & System Controls", "06. Admin Scope")

    tb_desc = slide7.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(4.5), Inches(5.5))
    tf = tb_desc.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "⚙️ System Administration Console"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = CRIMSON

    bullets = [
        "• Dynamic Metric Counters: Users (0), Patients (0), Caregivers (0), Admins (0) — calculated live from user entries.",
        "• Registered User Directory: Manage user roles and permission states.",
        "• Manual User Registration: Form to register Patient, Caregiver, or Admin accounts.",
        "• Account Controls: 1-click [ Block ], [ Unblock ], or [ Delete ] user actions.",
        "• Notification Gateway: Toggle Push, Email, and SMS notification channels."
    ]
    for b in bullets:
        p = tf.add_paragraph()
        p.text = b
        p.font.size = Pt(12)
        p.font.color.rgb = DARK_SLATE

    if os.path.exists(img_admin):
        slide7.shapes.add_picture(img_admin, Inches(5.5), Inches(1.4), width=Inches(7.0))

    # ==========================================
    # SLIDE 8: System Workflows & Architecture Planning
    # ==========================================
    slide8 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide8, "System Workflows & Architecture Planning", "07. Workflows")

    wf_cards = [
        ("🔐 1. Authentication Flow", "Role Category Selection ➔ Login Page ➔ Role Verification ➔ Redirect to /patient, /caregiver, or /admin."),
        ("⏰ 2. Reminder & Schedule Flow", "Add Medicine Form ➔ Daily Routine Setup (8am/1pm/8pm) ➔ Trigger Reminder Modal ➔ Log Intake (Taken/Missed/Snooze)."),
        ("📷 3. Prescription OCR Flow", "Upload Prescription Image ➔ AI/OCR Parsing ➔ Extract Name, Dosage, Stock ➔ Review & Save to Schedule."),
        ("🚨 4. Caregiver Escalation Flow", "Patient Misses Dose ➔ Log Missed Event ➔ Dispatch Alert to Caregiver Stream ➔ Caregiver Triggers Nudge / Call.")
    ]

    for i, (wtitle, wdesc) in enumerate(wf_cards):
        row = i // 2
        col = i % 2
        x = Inches(0.8 + col * 5.9)
        y = Inches(1.6 + row * 2.7)

        wbox = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.6), Inches(2.4))
        wbox.fill.solid()
        wbox.fill.fore_color.rgb = WHITE
        wbox.line.color.rgb = CARD_BORDER

        tb = slide8.shapes.add_textbox(x + Inches(0.15), y + Inches(0.15), Inches(5.3), Inches(2.1))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = wtitle
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = CRIMSON

        p2 = tf.add_paragraph()
        p2.text = wdesc
        p2.font.size = Pt(12)
        p2.font.color.rgb = DARK_SLATE

    # ==========================================
    # SLIDE 9: Verification, Testing & Git Repositories
    # ==========================================
    slide9 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide9, "Quality Assurance, Testing & Git Deliverables", "08. Verification")

    # Left: Testing Results
    t_box = slide9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    t_box.fill.solid()
    t_box.fill.fore_color.rgb = WHITE
    t_box.line.color.rgb = CARD_BORDER

    tb = slide9.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.2), Inches(4.8))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "🧪 Automated Test & Build Verification"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = CRIMSON

    results = [
        "✅ npm run format:check — Passed cleanly",
        "✅ npm run test — Passed cleanly",
        "✅ npm run build — 0 Errors (Production bundle generated)",
        "✅ Secret & Hygiene Check — Zero credentials committed",
        "✅ Responsive Layout Test — Verified across viewport sizes"
    ]
    for r in results:
        p = tf.add_paragraph()
        p.text = r
        p.font.size = Pt(13)
        p.font.color.rgb = DARK_SLATE

    # Right: Git Repositories Pushed
    g_box = slide9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2))
    g_box.fill.solid()
    g_box.fill.fore_color.rgb = ROSE_BG
    g_box.line.color.rgb = CARD_BORDER

    tb = slide9.shapes.add_textbox(Inches(7.0), Inches(1.8), Inches(5.3), Inches(4.8))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "🚀 GitHub Repositories Pushed"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = CRIMSON

    repos = [
        "1. Personal & Team Repository:\n   • URL: github.com/rajasrinallamilli/pillsync2\n   • Branches: main & rajasri",
        "2. Official Mentor Internship Repository:\n   • URL: github.com/GKSJ-Deepvision/PillSync\n   • Branch: intern/03-rajasri-nallamilli\n   • Fully compliant with INTERN_GUIDE.md roster rules."
    ]
    for rp in repos:
        p = tf.add_paragraph()
        p.text = rp
        p.font.size = Pt(12)
        p.font.color.rgb = DARK_SLATE

    # ==========================================
    # SLIDE 10: Conclusion & Milestone 2 Roadmap
    # ==========================================
    slide10 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide10, "Conclusion & Milestone 2 Roadmap", "09. Conclusion")

    c_box = slide10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.6), Inches(10.333), Inches(5.2))
    c_box.fill.solid()
    c_box.fill.fore_color.rgb = WHITE
    c_box.line.color.rgb = CARD_BORDER

    tb = slide10.shapes.add_textbox(Inches(1.8), Inches(1.9), Inches(9.7), Inches(4.6))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "🎉 Milestone 1 Successfully Completed!"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = CRIMSON

    summary_pts = [
        "• All assigned Milestone 1 frontend components, 25-screen wireframe blueprints, and system workflow diagrams are 100% complete and verified.",
        "• 🚀 Milestone 2 Roadmap Ahead (Weeks 3–4):",
        "   - Integrate frontend UI with Teammates' REST API endpoints & PostgreSQL database schema.",
        "   - Connect JWT Authentication tokens & Role-based Access Control (RBAC) middleware.",
        "   - Implement live push, email, and SMS reminder notifications.",
        "   - Enhance AI/OCR prescription scanning engine & refill predictions."
    ]
    for pt in summary_pts:
        p = tf.add_paragraph()
        p.text = pt
        p.font.size = Pt(14)
        p.font.color.rgb = DARK_SLATE

    p_end = tf.add_paragraph()
    p_end.text = "\nThank you! Questions & Feedback Welcome."
    p_end.font.size = Pt(16)
    p_end.font.bold = True
    p_end.font.color.rgb = CRIMSON
    p_end.alignment = PP_ALIGN.CENTER

    output_path = r"c:\Users\rajas\OneDrive\PillSync\PillSync_Milestone_1_Presentation.pptx"
    prs.save(output_path)
    print("PPTX generated successfully at:", output_path)

if __name__ == "__main__":
    create_presentation()
